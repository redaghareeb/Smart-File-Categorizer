"""
CTO Knowledge Asset Categorizer v6.0 — HYBRID EDITION
===================================================================
Features both a Blazing-Fast Keyword Engine AND a Deep-Semantic LLM Engine.
Automatically routes based on the user's GUI toggle.
"""

import os, sys, re, csv, json, math, shutil, logging, io, requests
from pathlib import Path
from collections import defaultdict
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed, TimeoutError

try:
    from tqdm import tqdm; HAS_TQDM = True
except ImportError:
    HAS_TQDM = False

import warnings
warnings.filterwarnings("ignore", message=".*Cannot set.*color.*")
warnings.filterwarnings("ignore", message=".*components specified.*")
for _ln in ("pdfminer","pdfminer.pdfpage","pdfminer.converter","pdfminer.cmapdb",
            "pdfminer.psparser","pdfminer.pdfdocument","pdfminer.pdfinterp"):
    logging.getLogger(_ln).setLevel(logging.CRITICAL)

log_fmt = logging.Formatter("%(asctime)s  %(levelname)-8s  %(message)s", datefmt="%H:%M:%S")
log = logging.getLogger("cto_v6"); log.setLevel(logging.INFO)
if sys.stderr is not None:
    _ch = logging.StreamHandler(); _ch.setFormatter(log_fmt); log.addHandler(_ch)
else:
    log.addHandler(logging.NullHandler())

# ═══════════ CONFIG ═══════════
class Config:
    max_workers=12; max_pages_pdf=20; max_rows_excel=500
    text_sample=10000; score_thr=3.5; prod_thr=8 
    min_distinct_kw=1   
    ocr_on=True; gpu=False; file_timeout=120; ocr_timeout=180 
    min_img_kb=20; max_img_mb=50; min_img_px=200; reset=False
    move_mode=False
    dedup=False
    dedup_action="report"
    transcribe_on=False
    whisper_model="medium"
    max_media_duration=300
    ffmpeg_path=""
    
    # 🧠 HYBRID AI TOGGLE
    use_llm = False
    
    # 🧠 OLLAMA CONFIGURATION
    ollama_url = "http://localhost:11434/api/generate"
    ollama_model = "qwen2.5:7b" 

    @classmethod
    def from_args(cls, args):
        c=cls()
        for a in args:
            if a=="--ocr": c.ocr_on=True
            elif a=="--gpu": c.gpu=True; c.ocr_on=True
            elif a.startswith("--timeout="): c.file_timeout=int(a.split("=")[1]); c.ocr_timeout=c.file_timeout*2
            elif a.startswith("--workers="): c.max_workers=int(a.split("=")[1])
            elif a=="--reset": c.reset=True
            elif a=="--move": c.move_mode=True
            elif a=="--dedup": c.dedup=True
            elif a=="--transcribe": c.transcribe_on=True
        c.max_workers=int(os.environ.get("MAX_WORKERS",c.max_workers))
        if os.environ.get("OCR_ENABLED")=="1": c.ocr_on=True
        if os.environ.get("GPU","").lower()=="true": c.gpu=True; c.ocr_on=True
        return c
CFG=Config()

# ═══════════ FILE TYPES ═══════════
DOC_EXT={".docx",".doc",".rtf",".odt",".txt",".md",".html",".htm",".xml",".json"}
PDF_EXT={".pdf"}
SHT_EXT={".xlsx",".xls",".xlsm",".csv",".tsv",".ods"}
PRS_EXT={".pptx",".ppt",".odp",".key"}
DIA_EXT={".vsd",".vsdx",".vsdm",".drawio"}
IMG_EXT={".png",".jpg",".jpeg",".tiff",".tif",".bmp",".webp"}
MEDIA_EXT=set()  
TXT_EXTS=DOC_EXT|PDF_EXT|SHT_EXT|PRS_EXT|DIA_EXT
ALL_EXTS=TXT_EXTS|IMG_EXT|MEDIA_EXT

# ═══════════ OCR ENGINE ═══════════
class SmartOCR:
    def __init__(self,gpu=False): self.name=None; self._r=None; self._gpu=gpu
    def _load(self):
        if self._r is not None or self.name=="__fail__": return
        try:
            import easyocr
            log.info("Loading EasyOCR ...")
            self._r=easyocr.Reader(['ar','en'],gpu=self._gpu,verbose=False)
            self.name="easyocr"; log.info("  ✅ EasyOCR (GPU:%s)",self._gpu); return
        except: pass
        try:
            from paddleocr import PaddleOCR
            self._r=PaddleOCR(lang='ar',use_angle_cls=True,use_gpu=self._gpu,show_log=False)
            self.name="paddleocr"; log.info("  ✅ PaddleOCR"); return
        except: pass
        try:
            import pytesseract; pytesseract.get_tesseract_version()
            self.name="tesseract"; log.info("  ✅ Tesseract"); return
        except: pass
        self.name="__fail__"; log.warning("⚠ No OCR Engine Found")
    def ocr(self,img):
        self._load()
        if not self.name or self.name=="__fail__": return ""
        try:
            from PIL import ImageOps,ImageFilter,Image as PILImage
            import numpy as np
            if img.mode=="RGBA":
                bg=PILImage.new("RGB",img.size,(255,255,255)); bg.paste(img,mask=img.split()[3]); img=bg
            elif img.mode not in("L","RGB"): img=img.convert("RGB")
            w,h=img.size
            if min(w,h)<800:
                s=max(2,1000//min(w,h)); img=img.resize((w*s,h*s),PILImage.LANCZOS)
            g=img.convert("L"); g=ImageOps.autocontrast(g,cutoff=1); g=g.filter(ImageFilter.SHARPEN)
            na=np.array(g)
            if self.name=="easyocr":
                res=self._r.readtext(na,paragraph=True,min_size=10,text_threshold=0.6)
                return "\n".join(r[1] for r in res if len(r)>=2 and (len(r)<3 or r[2]>0.2))
            elif self.name=="paddleocr":
                res=self._r.ocr(na,cls=True); lines=[]
                if res and res[0]:
                    for l in res[0]:
                        if l and len(l)>=2:
                            t=l[1][0] if isinstance(l[1],(list,tuple)) else str(l[1]); lines.append(t)
                return "\n".join(lines)
            elif self.name=="tesseract":
                import pytesseract; return pytesseract.image_to_string(g,lang="ara+eng").strip()
        except Exception as e: log.debug("OCR err: %s",e)
        return ""
    def ocr_file(self,fp):
        try:
            from PIL import Image; return self.ocr(Image.open(fp))
        except: return ""

_ocr=None
def get_ocr():
    global _ocr
    if _ocr is None: _ocr=SmartOCR(gpu=CFG.gpu)
    return _ocr

# ═══════════ TEXT EXTRACTION ═══════════
def _rd(fp):
    for e in("utf-8","utf-8-sig","utf-16","cp1256","iso-8859-6","cp1252","latin-1"):
        try:
            t=Path(fp).read_text(encoding=e,errors="ignore")
            if len(t.strip())>10: return t
        except: continue
    return ""

def ex_docx(fp):
    try:
        from docx import Document; d=Document(fp)
        p=[p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                c=[c.text.strip() for c in row.cells if c.text.strip()]
                if c: p.append(" | ".join(c))
        return "\n".join(p)
    except: return _rd(fp)

from contextlib import contextmanager
@contextmanager
def _suppress_stderr():
    old = sys.stderr
    try:
        sys.stderr = io.StringIO()
        yield
    finally:
        sys.stderr = old if old is not None else io.StringIO()

def ex_pdf(fp):
    parts=[]
    try:
        import pdfplumber
        with _suppress_stderr():
            with pdfplumber.open(fp) as pdf:
                for pg in pdf.pages[:CFG.max_pages_pdf]:
                    t=pg.extract_text()
                    if t: parts.append(t)
    except: pass
    c="\n".join(parts)
    if len(c.strip())<50: c=f"__NEEDS_OCR__ {c}"
    return c

def ex_xlsx(fp):
    try:
        from openpyxl import load_workbook; wb=load_workbook(fp,read_only=True,data_only=True)
        p=[]
        for ws in wb.worksheets:
            p.append(f"[{ws.title}]")
            for r in ws.iter_rows(max_row=CFG.max_rows_excel,values_only=True):
                v=[str(c) for c in r if c is not None]
                if v: p.append(" ".join(v))
        wb.close(); return "\n".join(p)
    except: return _rd(fp)

def ex_csv(fp):
    import csv as _c
    try:
        t=_rd(fp)
        if not t: return ""
        d="\t" if Path(fp).suffix.lower()==".tsv" else ","
        return "\n".join(" ".join(r) for i,r in enumerate(_c.reader(t.splitlines(),delimiter=d)) if i<800)
    except: return ""

def ex_pptx(fp):
    try:
        from pptx import Presentation; pr=Presentation(fp); p=[]
        for sl in pr.slides:
            for sh in sl.shapes:
                if sh.has_text_frame:
                    for pa in sh.text_frame.paragraphs:
                        if pa.text.strip(): p.append(pa.text)
                if sh.has_table:
                    for r in sh.table.rows:
                        c=[c.text.strip() for c in r.cells if c.text.strip()]
                        if c: p.append(" | ".join(c))
            if sl.has_notes_slide and sl.notes_slide.notes_text_frame:
                n=sl.notes_slide.notes_text_frame.text.strip()
                if n: p.append(n)
        return "\n".join(p)
    except: return _rd(fp)

EX={".docx":ex_docx,".doc":_rd,".rtf":_rd,".odt":_rd,".txt":_rd,".md":_rd,
    ".html":_rd,".htm":_rd,".xml":_rd,".json":_rd,".pdf":ex_pdf,
    ".xlsx":ex_xlsx,".xls":_rd,".xlsm":ex_xlsx,".csv":ex_csv,".tsv":ex_csv,
    ".ods":_rd,".pptx":ex_pptx,".ppt":_rd,".odp":_rd,".key":_rd,
    ".vsd":_rd,".vsdx":_rd,".vsdm":_rd,".drawio":_rd,
    ".png":lambda f:"__IMG__",".jpg":lambda f:"__IMG__",".jpeg":lambda f:"__IMG__",
    ".tiff":lambda f:"__IMG__",".tif":lambda f:"__IMG__",".bmp":lambda f:"__IMG__",".webp":lambda f:"__IMG__"}

def extract(fp):
    ext = Path(fp).suffix.lower()
    if ext in EX:
        return EX[ext](fp)[:CFG.text_sample]
    if ext in IMG_EXT:
        return "__IMG__"
    if ext in MEDIA_EXT:
        return "__MEDIA__"
    return _rd(fp)[:CFG.text_sample]

def transcribe_media(fp):
    try:
        import transcribe as tr
        if not tr.is_available(): return ""
        return tr.transcribe_text_only(fp, gpu=CFG.gpu, max_duration=CFG.max_media_duration, model_size=CFG.whisper_model)
    except Exception as e:
        log.debug("Transcription error for %s: %s", fp, e)
        return ""

def ocr_image(fp):
    try:
        sz=Path(fp).stat().st_size/1024
        if sz<CFG.min_img_kb or sz>CFG.max_img_mb*1024: return ""
        from PIL import Image; im=Image.open(fp); w,h=im.size
        if min(w,h)<CFG.min_img_px: im.close(); return ""
        im.close()
    except: return ""
    return get_ocr().ocr_file(fp)

def ocr_pdf(fp):
    try:
        import pdfplumber; ocr=get_ocr()
        if not ocr.name or ocr.name=="__fail__": return ""
        texts=[]
        with _suppress_stderr():
            with pdfplumber.open(fp) as pdf:
                for pg in pdf.pages[:10]:
                    t=ocr.ocr(pg.to_image(resolution=200).original)
                    if t and t.strip(): texts.append(t)
        return "\n".join(texts)
    except: return ""


# ═══════════ CATEGORIES & COMPILATION ═══════════
CATEGORIES = {}
COMPILED_CATS = {}

def _build_fn_pattern(hint):
    """Build regex for filename hint — loose substring match to bypass Arabic prefixes and brackets."""
    escaped = re.escape(hint.lower())
    return re.compile(escaped, re.IGNORECASE)

def _build_text_pattern(word):
    """Build regex for exact word boundary match in text content."""
    is_ar = bool(re.search(r'[\u0600-\u06FF]', word))
    escaped = re.escape(word.lower())
    if is_ar: return re.compile(r'(?:^|[^\u0600-\u06FF])(' + escaped + r')(?:[^\u0600-\u06FF]|$)', re.IGNORECASE)
    else: return re.compile(r'\b(' + escaped + r')\b', re.IGNORECASE)

def reload_categories(new_cats=None):
    global CATEGORIES, COMPILED_CATS
    if new_cats:
        CATEGORIES = new_cats
    else:
        try:
            from categories import load_categories
            CATEGORIES = load_categories()
        except ImportError: pass
        
    COMPILED_CATS = {}
    for cat, cfg in CATEGORIES.items():
        COMPILED_CATS[cat] = {
            "en": {k: (_build_text_pattern(k), w) for k, w in cfg.get("en", {}).items()},
            "ar": {k: (_build_text_pattern(k), w) for k, w in cfg.get("ar", {}).items()},
            "fh": [_build_fn_pattern(h) for h in cfg.get("fh", [])],
            "ph": [_build_fn_pattern(h) for h in cfg.get("ph", [])],
            "neg_en": [_build_text_pattern(k) for k in cfg.get("neg_en", [])],
            "neg_ar": [_build_text_pattern(k) for k in cfg.get("neg_ar", [])],
            "pp": cfg.get("pp", "UNKNOWN"),
            "personal": cfg.get("personal", False)
        }

# Initialize categories on module load
reload_categories()

# ═══════════ HYBRID ROUTING ENGINE ═══════════

def classify(text, filepath, filename):
    """Master Router: Sends document to AI LLM or Keyword Engine based on user toggle."""
    if CFG.use_llm:
        return classify_llm(text, filepath, filename)
    else:
        return classify_keyword(text, filepath, filename)

# --- ENGINE A: LLM (SEMANTIC) ---
def classify_llm(text, filepath, filename):
    # Create a mapping of category names to descriptive labels for the AI
    # This helps the LLM understand what's inside a folder like "01_Enterprise_Architecture"
    cat_hints = {
        "01_Enterprise_Architecture": "IT strategy, systems architecture, tech blueprints, and IT portals.",
        "20_Kids_Activities": "Worksheets, coloring pages, student games, and kids educational content.",
        "21_Personal_Documents": "Family records, IDs, CVs, school applications, and student exam results (Abnauna fi el-Kharej).",
        "00_General_Administration": "Official letters, memos, and general administrative circulars."
    }
    
    # Prepare the descriptive list for the prompt
    allowed_cats_with_desc = []
    for c in CATEGORIES.keys():
        desc = cat_hints.get(c, "Standard business or personal records matching this topic.")
        allowed_cats_with_desc.append(f"- {c}: {desc}")

    clean_text = text.replace('\n', ' ').strip()
    if len(clean_text) > 4000:
        clean_text = clean_text[:4000] + "... [TRUNCATED]"
        
    prompt = f"""
You are a highly intelligent document analyst. Your goal is to classify documents into the correct category based on their semantic meaning.

### CATEGORY DEFINITIONS:
{chr(10).join(allowed_cats_with_desc)}

### DOCUMENT TO ANALYZE:
Filename: {filename}
Text Sample: {clean_text if clean_text else "NO TEXT EXTRACTED. RELY ON FILENAME."}

### CLASSIFICATION RULES:
1. If the text mentions 'Student' (طالب), 'Primary Stage' (المرحلة الإبتدائية), or 'Our Children Abroad' (أبناؤنا في الخارج), it belongs in '21_Personal_Documents'.
2. Do NOT be confused by technical terms. A 'School Portal' is for students (Personal), NOT Enterprise Architecture (Corporate).
3. If the document is for a child or school-related, prioritize '20_Kids_Activities' or '21_Personal_Documents'.
4. Respond ONLY with pure JSON.

EXPECTED JSON:
{{
    "category": "The_Exact_Category_Name",
    "reason": "Explain your choice briefly."
}}
"""
    try:
        response = requests.post(CFG.ollama_url, json={
            "model": CFG.ollama_model, "prompt": prompt, "stream": False, "format": "json" 
        }, timeout=180)
        
        response.raise_for_status()
        output_text = response.json().get("response", "")
        json_match = re.search(r'\{.*\}', output_text, re.DOTALL)
        result_dict = json.loads(json_match.group(0)) if json_match else json.loads(output_text)
            
        selected_cat = result_dict.get("category", "99_Uncategorized")
        reason = result_dict.get("reason", "Categorized by LLM")
        
        if selected_cat not in CATEGORIES: selected_cat = "99_Uncategorized"
        pp = CATEGORIES.get(selected_cat, {}).get("pp", "UNKNOWN")
        return (selected_cat, 10.0, pp, [("🤖 " + reason, 10.0)])
        
    except Exception as e:
        log.debug(f"LLM Error for {filename}: {e}")
        return ("99_Uncategorized", 0.0, "UNKNOWN", [("LLM Error", 0.0)])

# --- ENGINE B: KEYWORDS (LEXICAL) ---
def classify_keyword(text, filepath, filename):
    t_lower = text.lower()
    fn_lower = filename.lower()
    ctx = (Path(filepath).parent.name + " " + fn_lower).lower()

    scores = {}
    lang = dlang(text)
    chk_en = lang in ("en", "mixed")
    chk_ar = lang in ("ar", "mixed")

    for cat_name, compiled in COMPILED_CATS.items():
        if cat_name.startswith("99_"): continue

        # Negative checks
        is_neg = False
        if chk_en:
            for pat in compiled["neg_en"]:
                if pat.search(t_lower): is_neg=True; break
        if chk_ar and not is_neg:
            for pat in compiled["neg_ar"]:
                if pat.search(t_lower): is_neg=True; break
        if is_neg: continue

        score = 0.0
        matches = []
        distinct_kw = set()

        # Filename hints
        for pattern in compiled["fh"]:
            if pattern.search(fn_lower):
                score += 4
                distinct_kw.add(f"__fn_{pattern.pattern}__")
                matches.append(("📄 filename", 4))

        # Path hints
        for pattern in compiled["ph"]:
            if pattern.search(ctx):
                score += 3
                distinct_kw.add(f"__ph_{pattern.pattern}__")
                matches.append(("📁 path", 3))

        # Content Matching (Logarithmic weight)
        if chk_en:
            for kw, (pat, w) in compiled["en"].items():
                hits = len(pat.findall(t_lower))
                if hits > 0:
                    pts = w * (1 + math.log(hits))
                    score += pts
                    distinct_kw.add(kw)
                    matches.append((kw, round(pts,2)))
        if chk_ar:
            for kw, (pat, w) in compiled["ar"].items():
                hits = len(pat.findall(t_lower))
                if hits > 0:
                    pts = w * (1 + math.log(hits))
                    score += pts
                    distinct_kw.add(kw)
                    matches.append((kw, round(pts,2)))

        if len(distinct_kw) >= CFG.min_distinct_kw:
            is_personal = compiled.get("personal", False)
            scores[cat_name] = {
                "score": score, 
                "matches": matches, 
                "pp": compiled.get("pp", "UNKNOWN"),
                "personal": is_personal
            }

    if not scores: return ("99_Uncategorized", 0.0, "NONE", [])

    # Personal Override Rule
    personal_cats = {k: v for k, v in scores.items() if v["personal"]}
    if personal_cats:
        best_p = max(personal_cats.items(), key=lambda x: x[1]["score"])
        if best_p[1]["score"] >= CFG.score_thr:
            m_sorted = sorted(best_p[1]["matches"], key=lambda x: -x[1])[:5]
            return (best_p[0], round(best_p[1]["score"],2), best_p[1]["pp"], m_sorted)

    best = max(scores.items(), key=lambda x: x[1]["score"])
    if best[1]["score"] >= CFG.score_thr:
        m_sorted = sorted(best[1]["matches"], key=lambda x: -x[1])[:5]
        return (best[0], round(best[1]["score"],2), best[1]["pp"], m_sorted)

    return ("99_Uncategorized", round(best[1]["score"],2), "NONE", [])


# ═══════════ HELPERS ═══════════
def hsz(n):
    for u in("B","KB","MB","GB"):
        if abs(n)<1024: return f"{n:.1f} {u}"
        n/=1024
    return f"{n:.1f} TB"

def uniq(d,fn):
    t=d/fn
    if not t.exists(): return t
    s,x=Path(fn).stem,Path(fn).suffix; c=1
    while True:
        t=d/f"{s}({c}){x}"
        if not t.exists(): return t
        c+=1

def dlang(t):
    if not t: return "unknown"
    ar=len(re.findall(r'[\u0600-\u06FF]',t[:5000]))
    en=len(re.findall(r'[a-zA-Z]',t[:5000]))
    tot=ar+en
    if tot==0: return "unknown"
    r=ar/tot
    return "ar" if r>0.7 else("mixed" if r>0.2 else "en")

def mkrow(src,fn,ext,sz,cat,sc,pp,kws,dst=None,dup=False,lang="unknown",ocr=False,err=None):
    return {"source":str(src),"destination":dst,"filename":fn,"extension":ext,"size":sz,
            "category":cat,"score":sc,"product_potential":pp,"top_keywords":kws,
            "duplicate":dup,"language":lang,"ocr_used":ocr,"error":err}

# ═══════════ PROCESS FILE ═══════════
def proc(src,dest,do_ocr=False,do_transcribe=False):
    fn=src.name; ext=src.suffix.lower()
    if ext not in ALL_EXTS: return None
    if not src.exists(): return None

    try:
        sz=src.stat().st_size
        if sz>200*1024*1024: return mkrow(src,fn,ext,sz,"99_Uncategorized",0,"NONE",[],err="Skipped >200MB")

        is_img=ext in IMG_EXT; is_scan=False; is_media=ext in MEDIA_EXT
        transcribed=False

        if is_img and not do_ocr: text=""
        elif is_media and not do_transcribe: text=""
        else: text=extract(str(src))

        if isinstance(text,str) and text.startswith("__NEEDS_OCR__"):
            is_scan=True; text=text.replace("__NEEDS_OCR__","").strip()
            if do_ocr:
                ot=ocr_pdf(str(src))
                if ot: text=ot
        if is_img and do_ocr and (not text or text=="__IMG__"):
            text=ocr_image(str(src))
        if text=="__IMG__": text=""

        if is_media and do_transcribe and (not text or text=="__MEDIA__"):
            text=transcribe_media(str(src))
            if text: transcribed=True
        if text=="__MEDIA__": text=""

        lang=dlang(text)
        cat,sc,pp,matches=classify(text,str(src),fn)
        cf=dest/cat; cf.mkdir(parents=True,exist_ok=True)
        df=uniq(cf,fn); dup=df.name!=fn

        shutil.copy2(src,df)

        if CFG.move_mode:
            if df.exists() and df.stat().st_size == sz:
                try: src.unlink()
                except Exception as e: log.debug("Could not delete source %s: %s", src, e)

        return mkrow(src,fn,ext,sz,cat,sc,pp,[m[0] for m in matches],
                    dst=str(df),dup=dup,lang=lang,
                    ocr=(is_img and do_ocr)or(is_scan and do_ocr) or transcribed)
    except Exception as e:
        return mkrow(src,fn,ext,0,"ERROR",0,"NONE",[],err=str(e))

# ═══════════ BATCH PROCESSOR ═══════════
_progress_cb = None    
_log_cb = None         
_cancel = None         

def set_callbacks(progress_fn=None, log_fn=None, cancel_event=None):
    global _progress_cb, _log_cb, _cancel
    _progress_cb = progress_fn
    _log_cb = log_fn
    _cancel = cancel_event

def batch(files,dest,done,rpath,do_ocr,timeout,desc,workers=None):
    results=[]; errs=0; tos=0
    # Throttling workers for LLM to prevent memory crashes
    w = min(4, CFG.max_workers) if CFG.use_llm else (workers or CFG.max_workers)
    
    prog=tqdm(total=len(files),desc=desc,unit="ملف") if (HAS_TQDM and sys.stderr is not None) else None
    save_n=max(200,len(files)//10)
    total=len(files)
    with ThreadPoolExecutor(max_workers=w) as ex:
        futs={ex.submit(proc,f,dest,do_ocr):f for f in files}
        dc=0
        for fut in as_completed(futs):
            if _cancel and _cancel.is_set():
                ex.shutdown(wait=False,cancel_futures=True)
                if _log_cb: _log_cb("⚠ User Canceled.")
                break
            try:
                r=fut.result(timeout=timeout)
                if r:
                    results.append(r)
                    if r["error"]: errs+=1
                    done.add(r["source"])
                    if _progress_cb: _progress_cb(dc+1,total,r.get("filename",""),r.get("category",""))
            except TimeoutError:
                fp=futs[fut]; tos+=1
                results.append(mkrow(fp,fp.name,fp.suffix.lower(),0,"99_Uncategorized",0,"NONE",[],err=f"Timeout({timeout}s)"))
                done.add(str(fp))
            except Exception as e:
                fp=futs[fut]
                results.append(mkrow(fp,fp.name,fp.suffix.lower(),0,"ERROR",0,"NONE",[],err=str(e)))
                done.add(str(fp))
            dc+=1
            if prog: prog.update(1)
            elif dc%1000==0: log.info("  ... %d / %d",dc,total)
            if dc%save_n==0:
                try: rpath.write_text(json.dumps(list(done),ensure_ascii=False),encoding="utf-8")
                except: pass
    if prog: prog.close()
    try: rpath.write_text(json.dumps(list(done),ensure_ascii=False),encoding="utf-8")
    except: pass
    log.info("  ✅ %d processed | %d errors | %d timeouts",len(results),errs,tos)
    return results

import hashlib

def _file_hash_partial(filepath, chunk_size=8192):
    try:
        p = Path(filepath); sz = p.stat().st_size; h = hashlib.md5(); h.update(str(sz).encode())
        with open(filepath, 'rb') as f:
            h.update(f.read(chunk_size))
            if sz > chunk_size * 2: f.seek(-chunk_size, 2); h.update(f.read(chunk_size))
        return h.hexdigest()
    except Exception: return None

def _file_hash_full(filepath):
    try:
        h = hashlib.md5()
        with open(filepath, 'rb') as f:
            for chunk in iter(lambda: f.read(65536), b''): h.update(chunk)
        return h.hexdigest()
    except Exception: return None

def _extract_metadata(filepath):
    meta = {}; ext = Path(filepath).suffix.lower()
    try:
        if ext == '.docx':
            from docx import Document; cp = Document(filepath).core_properties
            meta['title'] = cp.title or ''; meta['author'] = cp.author or ''
        elif ext == '.pptx':
            from pptx import Presentation; cp = Presentation(filepath).core_properties
            meta['title'] = cp.title or ''; meta['author'] = cp.author or ''
        elif ext == '.xlsx':
            from openpyxl import load_workbook; wb = load_workbook(filepath, read_only=True); props = wb.properties
            meta['title'] = props.title or ''; meta['author'] = props.creator or ''; wb.close()
        elif ext == '.pdf':
            import pdfplumber
            with _suppress_stderr():
                with pdfplumber.open(filepath) as pdf:
                    info = pdf.metadata or {}; meta['title'] = info.get('Title', ''); meta['author'] = info.get('Author', '')
    except Exception: pass
    return meta

def _pick_best_copy(group):
    def sort_key(item):
        try: mtime = Path(item['destination']).stat().st_mtime if item.get('destination') else 0
        except Exception: mtime = 0
        path_len = len(item.get('destination', '')) if item.get('destination') else 9999
        return (-item.get('score', 0), -mtime, path_len)
    return sorted(group, key=sort_key)[0] 

def find_duplicates(results, dest_path):
    if _log_cb: _log_cb("🔍 Scanning for duplicates ...")
    valid = [r for r in results if r.get('destination') and not r.get('error') and Path(r['destination']).exists()]
    if not valid: return [], {}

    size_groups = defaultdict(list)
    for r in valid:
        try:
            sz = Path(r['destination']).stat().st_size
            if sz > 0: size_groups[sz].append(r)
        except Exception: pass

    candidates = {sz: items for sz, items in size_groups.items() if len(items) >= 2}
    if not candidates: return [], {"total_checked": len(valid), "duplicates_found": 0}

    partial_groups = defaultdict(list)
    for sz, items in candidates.items():
        for r in items:
            ph = _file_hash_partial(r['destination'])
            if ph: partial_groups[ph].append(r)

    partial_dupes = {h: items for h, items in partial_groups.items() if len(items) >= 2}
    if not partial_dupes: return [], {"total_checked": len(valid), "duplicates_found": 0}

    duplicate_groups = []; total_dupes = 0
    for ph, items in partial_dupes.items():
        full_hash_groups = defaultdict(list)
        for r in items:
            fh = _file_hash_full(r['destination'])
            if fh: full_hash_groups[fh].append(r)

        for fh, group in full_hash_groups.items():
            if len(group) >= 2:
                best = _pick_best_copy(group); removes = [r for r in group if r is not best]; total_dupes += len(removes)
                match_type = 'exact_content' if len(set(r['filename'] for r in group)) == 1 else 'same_content_diff_name'
                duplicate_groups.append({'hash': fh, 'keep': best, 'remove': removes, 'match_type': match_type, 'size': group[0].get('size', 0), 'count': len(group)})

    meta_groups = defaultdict(list)
    for r in valid:
        dest = r.get('destination')
        if not dest: continue
        ext = Path(dest).suffix.lower()
        if ext in ('.docx', '.pptx', '.xlsx', '.pdf'):
            meta = _extract_metadata(dest)
            title = (meta.get('title', '') or '').strip().lower(); author = (meta.get('author', '') or '').strip().lower()
            if title and len(title) > 3:
                meta_groups[f"{title}|{author}|{ext}"].append({**r, '_meta': meta})

    near_dupes = []
    for key, group in meta_groups.items():
        if len(group) >= 2:
            dests = set(r['destination'] for r in group); already_found = False
            for dg in duplicate_groups:
                existing_dests = {dg['keep']['destination']} | {r['destination'] for r in dg['remove']}
                if dests & existing_dests: already_found = True; break
            if not already_found:
                best = _pick_best_copy(group); removes = [r for r in group if r is not best]
                near_dupes.append({'hash': f"meta:{key[:50]}", 'keep': best, 'remove': removes, 'match_type': 'metadata_match', 'size': group[0].get('size', 0), 'count': len(group), 'meta_key': key})

    all_groups = duplicate_groups + near_dupes
    total_removable = sum(len(g['remove']) for g in all_groups)
    total_space = sum(sum(Path(r['destination']).stat().st_size for r in g['remove'] if r.get('destination') and Path(r['destination']).exists()) for g in all_groups)

    report = {
        "total_checked": len(valid), "exact_duplicate_groups": len(duplicate_groups), "near_duplicate_groups": len(near_dupes),
        "duplicates_found": total_removable, "space_recoverable": total_space, "space_recoverable_human": hsz(total_space),
    }
    return all_groups, report


def apply_dedup(duplicate_groups, dest_path, action="move_to_dupes"):
    if action == "report": return 0
    import stat
    dupes_folder = dest_path / "_duplicates"; acted = 0
    
    for group in duplicate_groups:
        for r in group['remove']:
            fp = Path(r['destination'])
            if not fp.exists(): continue
            try:
                if action == "delete":
                    fp.chmod(stat.S_IWRITE); fp.unlink(); acted += 1
                elif action == "move_to_dupes":
                    dupes_folder.mkdir(parents=True, exist_ok=True)
                    sub = dupes_folder / group.get('match_type', 'unknown'); sub.mkdir(exist_ok=True)
                    target = sub / fp.name
                    if target.exists():
                        stem, suffix = fp.stem, fp.suffix; c = 1
                        while target.exists(): target = sub / f"{stem}({c}){suffix}"; c += 1
                    shutil.move(str(fp), str(target)); acted += 1
            except Exception as e: log.debug("Dedup action failed for %s: %s", fp, e)
    return acted

def save_dedup_report(duplicate_groups, dedup_stats, dest_path):
    csv_path = dest_path / "_DUPLICATES_REPORT.csv"
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=["group_id", "match_type", "action", "filename", "category", "score", "size", "path", "title", "author", "modified"])
        w.writeheader()
        for i, group in enumerate(duplicate_groups, 1):
            keep = group['keep']; meta = _extract_metadata(keep['destination']) if keep.get('destination') and Path(keep['destination']).exists() else {}
            w.writerow({"group_id": i, "match_type": group['match_type'], "action": "✅ KEEP", "filename": keep.get('filename', ''), "category": keep.get('category', ''), "score": keep.get('score', 0), "size": keep.get('size', 0), "path": keep.get('destination', ''), "title": meta.get('title', ''), "author": meta.get('author', ''), "modified": meta.get('fs_modified', '')})
            for r in group['remove']:
                meta_r = _extract_metadata(r['destination']) if r.get('destination') and Path(r['destination']).exists() else {}
                w.writerow({"group_id": i, "match_type": group['match_type'], "action": "❌ DUPLICATE", "filename": r.get('filename', ''), "category": r.get('category', ''), "score": r.get('score', 0), "size": r.get('size', 0), "path": r.get('destination', r.get('source', '')), "title": meta_r.get('title', ''), "author": meta_r.get('author', ''), "modified": meta_r.get('fs_modified', '')})
    
    json_path = dest_path / "_DUPLICATES_REPORT.json"
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump({"timestamp": datetime.now().isoformat(), "stats": dedup_stats, "groups": [{"match_type": g['match_type'], "hash": g['hash'], "file_count": g['count'], "keep": g['keep']['destination'], "remove": [r['destination'] for r in g['remove']]} for g in duplicate_groups]}, f, indent=2, ensure_ascii=False)

def reports(results,dest,src):
    stats=defaultdict(lambda:{"c":0,"s":0,"pp":"","l":defaultdict(int),"o":0}); prods=[]
    for r in results:
        cat=r["category"]; stats[cat]["c"]+=1; stats[cat]["s"]+=r.get("size",0)
        stats[cat]["pp"]=r.get("product_potential",""); stats[cat]["l"][r.get("language","unknown")]+=1
        if r.get("ocr_used"): stats[cat]["o"]+=1
        if r.get("product_potential") in("HIGH","VERY_HIGH") and r.get("score",0)>=CFG.prod_thr: prods.append(r)
    
    rj=dest/"_categorization_report.json"
    with open(rj,"w",encoding="utf-8") as f:
        json.dump({"timestamp":datetime.now().isoformat(),"source":str(src),"total":len(results),
                   "category_summary":{c:{"count":s["c"],"size":hsz(s["s"]),"product_potential":s.get("pp","")} for c,s in sorted(stats.items())},
                   "product_candidates":len(prods),"files":results},f,indent=2,ensure_ascii=False)

    rc=dest/"_categorization_report.csv"
    with open(rc,"w",newline="",encoding="utf-8-sig") as f:
        w=csv.DictWriter(f,fieldnames=["filename","extension","language","ocr_used","category","score","product_potential","top_keywords","duplicate","size","source","destination","error"])
        w.writeheader()
        for r in results:
            row=dict(r); row["top_keywords"]=" | ".join(str(x) for x in row.get("top_keywords",[])); w.writerow(row)

def run(source,destination):
    if isinstance(source, (list, tuple)): sources = [Path(s).resolve() for s in source if s and Path(s).is_dir()]
    else:
        sp = Path(source).resolve()
        if not sp.is_dir(): return
        sources = [sp]

    if not sources: return
    dp=Path(destination).resolve(); dp.mkdir(parents=True,exist_ok=True)

    log_path = dp / "_categorization.log"
    existing = [h for h in log.handlers if isinstance(h, logging.FileHandler) and hasattr(h, 'baseFilename') and Path(h.baseFilename).name == "_categorization.log"]
    if not existing:
        fh=logging.FileHandler(log_path, encoding="utf-8"); fh.setFormatter(log_fmt); log.addHandler(fh)

    if _progress_cb: _progress_cb(0, 0, "Scanning files...", "")
    rp=dp/"_processed_files.json"; done=set()
    if not CFG.reset and rp.exists():
        try: done=set(json.loads(rp.read_text(encoding="utf-8")))
        except: pass

    tp = []; ip = []; mp = []; discovered_count = 0
    for sp in sources:
        for root, _, fnames in os.walk(sp):
            if _cancel and _cancel.is_set(): return
            root_path = Path(root)
            for fn in fnames:
                ext = fn[fn.rfind('.'):].lower() if '.' in fn else ""
                if ext in ALL_EXTS:
                    full_path_str = str(root_path / fn)
                    if full_path_str not in done:
                        p = root_path / fn
                        if ext in TXT_EXTS: tp.append(p)
                        elif ext in IMG_EXT: ip.append(p)
                        elif ext in MEDIA_EXT: mp.append(p)
                        discovered_count += 1
                        if discovered_count % 1000 == 0 and _progress_cb: _progress_cb(0, 0, f"Found {discovered_count:,} new files...", "")

    if not tp and not ip and not mp: return
    results=[]

    if tp:
        results+=batch(tp,dp,done,rp,False,CFG.file_timeout,"📄 Documents")
        if _cancel and _cancel.is_set(): return

    if ip:
        if CFG.ocr_on:
            o=get_ocr(); o._load()
            if not o.name or o.name=="__fail__": CFG.ocr_on=False
        w = min(4, CFG.max_workers) if (CFG.ocr_on and not CFG.gpu) else CFG.max_workers
        t=CFG.ocr_timeout if CFG.ocr_on else CFG.file_timeout
        results+=batch(ip,dp,done,rp,CFG.ocr_on,t,"🖼️ Images",workers=w)
        if _cancel and _cancel.is_set(): return

    if mp:
        if CFG.transcribe_on:
            try:
                import transcribe as tr
                if not tr.is_available(): CFG.transcribe_on=False
            except ImportError: CFG.transcribe_on=False
        
        media_workers = 1 if CFG.transcribe_on else CFG.max_workers
        media_results = batch_media(mp,dp,done,rp,CFG.transcribe_on,
                                     CFG.file_timeout*3 if CFG.transcribe_on else CFG.file_timeout,
                                     "🎵 Audio/Video", workers=media_workers)
        results += media_results
        if _cancel and _cancel.is_set(): return

    reports(results,dp,sources[0] if len(sources)==1 else Path(destination))

    if CFG.dedup and results:
        dup_groups, dup_stats = find_duplicates(results, dp)
        if dup_groups:
            save_dedup_report(dup_groups, dup_stats, dp)
            if CFG.dedup_action != "report": apply_dedup(dup_groups, dp, CFG.dedup_action)

def batch_media(files,dest,done,rpath,do_transcribe,timeout,desc,workers=None):
    results=[]; errs=0; tos=0; w=workers or CFG.max_workers
    prog=tqdm(total=len(files),desc=desc,unit="ملف") if (HAS_TQDM and sys.stderr is not None) else None
    save_n=max(50,len(files)//10); total=len(files)
    with ThreadPoolExecutor(max_workers=w) as ex:
        futs={ex.submit(proc,f,dest,False,do_transcribe):f for f in files}
        dc=0
        for fut in as_completed(futs):
            if _cancel and _cancel.is_set(): ex.shutdown(wait=False,cancel_futures=True); break
            try:
                r=fut.result(timeout=timeout)
                if r:
                    results.append(r); done.add(r["source"])
                    if r["error"]: errs+=1
                    if _progress_cb: _progress_cb(dc+1,total,r.get("filename",""),r.get("category",""))
            except TimeoutError:
                fp=futs[fut]; tos+=1; results.append(mkrow(fp,fp.name,fp.suffix.lower(),0,"99_Uncategorized",0,"NONE",[],err=f"Timeout({timeout}s)")); done.add(str(fp))
            except Exception as e:
                fp=futs[fut]; results.append(mkrow(fp,fp.name,fp.suffix.lower(),0,"ERROR",0,"NONE",[],err=str(e))); done.add(str(fp))
            dc+=1
            if prog: prog.update(1)
            if dc%save_n==0:
                try: rpath.write_text(json.dumps(list(done),ensure_ascii=False),encoding="utf-8")
                except: pass
    if prog: prog.close()
    try: rpath.write_text(json.dumps(list(done),ensure_ascii=False),encoding="utf-8")
    except: pass
    return results